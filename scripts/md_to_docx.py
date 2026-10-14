"""Render meeting minutes Markdown to the standard Word layout.

Usage: md_to_docx.py <minutes.md> [out.docx]

The layout is the company's standard meeting-minutes format:
A4, 1-inch margins, Microsoft JhengHei 11pt; title lines centred, bold, navy 17pt;
「一、」 headings bold navy 14pt; 「（一）」 headings bold green 12pt; bullets and numbered
lists with line 320 / after 60; body paragraphs line 340 / after 100; tables 9026 dxa wide,
grey borders, navy header with white bold text; centred page number in the footer.
Markdown supported: # / ## / ### headings, paragraphs, **bold**, bullet and numbered lists, pipe tables,
[links](url) (kept as clickable links) and ![images](src) (data: URIs as Google's Markdown export writes them,
http(s) URLs, or files next to the minutes; scaled down to the text width).
"""
import base64
import io
import sys
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from markdown_it import MarkdownIt

FONT = 'Microsoft JhengHei'
NAVY = RGBColor(0x14, 0x30, 0x4A)
GREEN = RGBColor(0x1D, 0x6F, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINK = RGBColor(0x11, 0x55, 0xCC)
MAX_IMAGE = Cm(15.9)                 # A4 minus the two 2.54 cm margins
BASE_DIR = [Path('.')]               # folder of the minutes being built, for relative image paths
# tag -> (size pt, colour, space before twips, space after twips, centred)
HEADING = {'h1': (17, NAVY, 0, 80, True), 'h2': (14, NAVY, 280, 120, False), 'h3': (12, GREEN, 200, 80, False)}
CONTENT_DXA = 9026
# header text -> column width in dxa; columns not listed share what is left
WIDTH_DXA = {'編號': 700, '負責': 1900, '負責單位': 1900, '確認對象': 2000,
             '說明與執行方向': 4126, '說明': 4626}


def set_font(run, size=11, bold=None, color=None):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def spacing(par, before=0, after=0, line=None):
    fmt = par.paragraph_format
    fmt.space_before = Twips(before)
    fmt.space_after = Twips(after)
    if line:
        fmt.line_spacing = line / 240


def add_inline(par, token, size=11, color=None, bold_all=False):
    bold = bold_all
    last = ''
    href = None
    for child in token.children or []:
        if child.type == 'strong_open':
            bold = True
        elif child.type == 'strong_close':
            bold = bold_all
        elif child.type == 'link_open':
            href = child.attrGet('href')
        elif child.type == 'link_close':
            href = None
        elif child.type in ('text', 'code_inline'):
            if href:
                add_link(par, href, child.content, size, bold)
            else:
                set_font(par.add_run(child.content), size, bold, color)
            last = child.content
        elif child.type == 'image':
            add_image(par, child.attrGet('src'), child.content)
        elif child.type == 'softbreak':
            # a source line wrap, not a real break: nothing between CJK characters, a space between words
            if last[-1:].isascii() and last[-1:].isalnum():
                set_font(par.add_run(' '), size, bold, color)
        elif child.type == 'hardbreak':
            par.add_run().add_break()


def add_link(par, url, text, size, bold):
    run = par.add_run(text)
    set_font(run, size, bold, LINK)
    run.font.underline = True
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    par._p.remove(run._r)
    link.append(run._r)
    par._p.append(link)


def image_bytes(src: str) -> bytes:
    if src.startswith('data:'):
        return base64.b64decode(src.split(',', 1)[1])
    if src.startswith(('http://', 'https://')):
        return urllib.request.urlopen(src, timeout=30).read()
    return (BASE_DIR[0] / urllib.parse.unquote(src)).read_bytes()   # markdown-it percent-encodes 中文 and spaces


def add_image(par, src, alt=''):
    try:
        data = image_bytes(src)
        width = min(Image.from_blob(data).width, MAX_IMAGE)
        par.add_run().add_picture(io.BytesIO(data), width=width)
    except Exception as err:                       # one bad picture must not stop the whole minutes
        print(f'WARNING: image not placed ({err.__class__.__name__}): {src[:80]}', file=sys.stderr)
        set_font(par.add_run(f'［圖片無法載入：{alt or src[:60]}］'), 11, False, LINK)


def element(tag, **attrs):
    el = OxmlElement(tag.replace('_', ':', 1))
    for key, val in attrs.items():
        el.set(qn(key.replace('_', ':', 1)), str(val))
    return el


def style_cell(cell, width, header):
    pr = cell._element.get_or_add_tcPr()
    for old in pr.findall(qn('w:tcW')):
        pr.remove(old)
    pr.append(element('w_tcW', w_type='dxa', w_w=width))
    borders = element('w_tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        borders.append(element(f'w_{side}', w_val='single', w_color='A0A8B0', w_sz=4))
    pr.append(borders)
    if header:
        pr.append(element('w_shd', w_fill='14304A', w_color='auto', w_val='clear'))
    margins = element('w_tcMar')
    for side, val in (('top', 80), ('left', 100), ('bottom', 80), ('right', 100)):
        margins.append(element(f'w_{side}', w_type='dxa', w_w=val))
    pr.append(margins)


def column_widths(header_cells):
    names = [''.join(c.content for c in (h.children or [])) for h in header_cells]
    fixed = [WIDTH_DXA.get(n) for n in names]
    free = fixed.count(None)
    used = sum(w for w in fixed if w)
    if free and CONTENT_DXA - used < 1200 * free:          # fixed columns leave too little: shrink them
        scale = (CONTENT_DXA - 1200 * free) / used
        fixed = [int(w * scale) if w else None for w in fixed]
        used = sum(w for w in fixed if w)
    rest = (CONTENT_DXA - used) // free if free else 0
    widths = [w or rest for w in fixed]
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_table(doc, rows):
    widths = column_widths(rows[0][1])
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.autofit = False
    tbl_pr = table._element.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = element('w_tblW')
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(CONTENT_DXA))
    for col, w in zip(table._element.tblGrid.findall(qn('w:gridCol')), widths):
        col.set(qn('w:w'), str(w))
    for r_idx, (is_head, cells) in enumerate(rows):
        for c_idx, inline in enumerate(cells[:len(widths)]):
            cell = table.cell(r_idx, c_idx)
            style_cell(cell, widths[c_idx], is_head)
            par = cell.paragraphs[0]
            spacing(par, line=300)
            add_inline(par, inline, 11, WHITE if is_head else None, bold_all=is_head)


def list_style(base: str, depth: int) -> str:
    return base if depth == 0 else f'{base} {min(depth + 1, 3)}'


def restart_numbering(doc, style_name: str, start: int) -> int:
    """Give each numbered list its own counter so a later list starts at 1 again."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles[style_name].element.pPr.numPr.numId.val
    abstract = next(n.abstractNumId.val for n in numbering.num_lst if n.numId == style_num)
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(start)
    return num.numId


def page_number_footer(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    run._element.append(element('w_fldChar', w_fldCharType='begin'))
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    run._element.append(instr)
    run._element.append(element('w_fldChar', w_fldCharType='end'))
    set_font(run, 9)


def build(md_path: Path, out_path: Path):
    BASE_DIR[0] = md_path.parent
    tokens = MarkdownIt('commonmark').enable('table').parse(md_path.read_text(encoding='utf-8'))
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    for level, (size, color, *_rest) in zip((1, 2, 3), HEADING.values()):
        style = doc.styles[f'Heading {level}']
        style.font.name, style.font.size, style.font.bold, style.font.color.rgb = FONT, Pt(size), True, color
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Twips(1440))
    page_number_footer(sec)

    lists = []                       # stack of (style name, numId or None) per open list
    table_rows, row, in_head = None, None, False
    for i, t in enumerate(tokens):
        prev = tokens[i - 1].type if i else ''
        if t.type == 'inline' and prev == 'heading_open':
            tag = tokens[i - 1].tag
            size, color, before, after, centred = HEADING.get(tag, HEADING['h3'])
            # real Heading styles, so Google Docs keeps them as headings and its Markdown export writes # again
            par = doc.add_paragraph(style=f'Heading {min(int(tag[1]), 3)}')
            if centred:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(par, before, after)
            add_inline(par, t, size, color, bold_all=True)
        elif t.type == 'bullet_list_open':
            lists.append((list_style('List Bullet', len(lists)), None))
        elif t.type == 'ordered_list_open':
            style = list_style('List Number', len(lists))
            lists.append((style, restart_numbering(doc, style, int(t.attrGet('start') or 1))))
        elif t.type in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif t.type == 'table_open':
            table_rows = []
        elif t.type == 'thead_open':
            in_head = True
        elif t.type == 'thead_close':
            in_head = False
        elif t.type == 'tr_open':
            row = []
        elif t.type == 'tr_close':
            table_rows.append((in_head, row))
        elif t.type == 'inline' and prev in ('th_open', 'td_open'):
            row.append(t)
        elif t.type == 'table_close':
            add_table(doc, table_rows)
            table_rows, row = None, None
        elif t.type == 'inline' and prev == 'paragraph_open':
            if lists:
                style, num_id = lists[-1]
                par = doc.add_paragraph(style=style)
                if num_id:
                    par._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
                spacing(par, 0, 60, 320)
            else:
                par = doc.add_paragraph()
                spacing(par, 0, 100, 340)
            add_inline(par, t, 11)
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    src = Path(sys.argv[1])
    print(build(src, Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('.docx')))
